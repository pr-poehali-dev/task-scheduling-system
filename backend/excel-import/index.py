"""
Импорт Excel-файлов МБУ АПХ Абакан.
Принимает base64-файл + тип (employees | work_types | timesheet | orders).
Парсит все листы, сохраняет в БД, возвращает статистику.
"""
import json
import base64
import os
import psycopg2
from parsers import parse_employees, parse_work_types, parse_timesheet, parse_orders


def get_conn():
    return psycopg2.connect(os.environ["DATABASE_URL"])


CORS = {
    "Access-Control-Allow-Origin": "*",
    "Access-Control-Allow-Methods": "GET, POST, OPTIONS",
    "Access-Control-Allow-Headers": "Content-Type",
}


def handler(event: dict, context) -> dict:
    """Импорт Excel-файлов: employees, work_types, timesheet, orders"""
    if event.get("httpMethod") == "OPTIONS":
        return {"statusCode": 200, "headers": CORS, "body": ""}

    if event.get("httpMethod") == "GET":
        path = event.get("path", "")
        if "/log" in path:
            return get_import_log(event)
        return get_data(event)

    body = json.loads(event.get("body") or "{}")
    file_b64 = body.get("file")
    file_type = body.get("type")  # employees | work_types | timesheet | orders
    file_name = body.get("name", "file.xlsx")

    if not file_b64 or not file_type:
        return {
            "statusCode": 400,
            "headers": CORS,
            "body": json.dumps({"error": "Нужны поля file (base64) и type"}),
        }

    file_bytes = base64.b64decode(file_b64)

    # Диагностический режим — возвращает структуру файла без сохранения
    if file_type == "debug":
        return debug_file(file_bytes)

    conn = get_conn()
    cur = conn.cursor()
    rows_ok = 0
    rows_err = 0
    err_msg = ""

    try:
        if file_type == "employees":
            rows_ok, rows_err, err_msg = import_employees(cur, file_bytes)
        elif file_type == "work_types":
            rows_ok, rows_err, err_msg = import_work_types(cur, file_bytes)
        elif file_type == "timesheet":
            rows_ok, rows_err, err_msg = import_timesheet(cur, file_bytes)
        elif file_type == "orders":
            rows_ok, rows_err, err_msg = import_orders(cur, file_bytes)
        else:
            return {
                "statusCode": 400,
                "headers": CORS,
                "body": json.dumps({"error": f"Неизвестный тип: {file_type}"}),
            }

        conn.commit()

        # Лог импорта
        cur.execute(
            "INSERT INTO import_log (file_name, file_type, rows_imported, rows_error, status, error_msg) "
            "VALUES (%s, %s, %s, %s, %s, %s)",
            (file_name, file_type, rows_ok, rows_err, "ok" if not err_msg else "partial", err_msg[:500] if err_msg else "")
        )
        conn.commit()

    except Exception as e:
        conn.rollback()
        cur.execute(
            "INSERT INTO import_log (file_name, file_type, rows_imported, rows_error, status, error_msg) "
            "VALUES (%s, %s, %s, %s, %s, %s)",
            (file_name, file_type, 0, 0, "error", str(e)[:500])
        )
        conn.commit()
        cur.close()
        conn.close()
        return {
            "statusCode": 500,
            "headers": CORS,
            "body": json.dumps({"error": str(e), "imported": 0}),
        }

    cur.close()
    conn.close()

    return {
        "statusCode": 200,
        "headers": CORS,
        "body": json.dumps({
            "ok": True,
            "type": file_type,
            "imported": rows_ok,
            "errors": rows_err,
            "warning": err_msg or None,
        }),
    }


def import_employees(cur, file_bytes: bytes):
    records = parse_employees(file_bytes)
    ok, err = 0, 0
    errors = []
    for r in records:
        try:
            cur.execute("""
                INSERT INTO employees (tab_number, full_name, position, department, category, work_schedule, status)
                VALUES (%s, %s, %s, %s, %s, %s, 'active')
                ON CONFLICT (tab_number) DO UPDATE SET
                    full_name = EXCLUDED.full_name,
                    position = EXCLUDED.position,
                    category = EXCLUDED.category,
                    work_schedule = EXCLUDED.work_schedule,
                    updated_at = NOW()
            """, (r["tab_number"], r["full_name"], r["position"], r["department"], r["category"], r["work_schedule"]))
            ok += 1
        except Exception as e:
            err += 1
            errors.append(str(e)[:100])
    return ok, err, "; ".join(errors[:5])


def import_work_types(cur, file_bytes: bytes):
    records = parse_work_types(file_bytes)
    ok, err = 0, 0
    errors = []
    for r in records:
        try:
            cur.execute("""
                INSERT INTO work_types (code, name, group_name, unit, norm_per_hour)
                VALUES (%s, %s, %s, %s, %s)
            """, (r["code"], r["name"], r["group_name"], r["unit"], r["norm_per_hour"]))
            ok += 1
        except Exception as e:
            err += 1
            errors.append(str(e)[:100])
    return ok, err, "; ".join(errors[:5])


def import_timesheet(cur, file_bytes: bytes):
    records = parse_timesheet(file_bytes)
    ok, err = 0, 0
    errors = []
    for r in records:
        try:
            # Пытаемся найти employee_id по табельному номеру
            emp_id = None
            if r["tab_number"]:
                cur.execute("SELECT id FROM employees WHERE tab_number = %s LIMIT 1", (r["tab_number"],))
                row = cur.fetchone()
                if row:
                    emp_id = row[0]

            cur.execute("""
                INSERT INTO timesheet (employee_id, employee_name, tab_number, year, month, day, mark, hours)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
                ON CONFLICT (tab_number, year, month, day) DO UPDATE SET
                    mark = EXCLUDED.mark,
                    hours = EXCLUDED.hours,
                    employee_id = COALESCE(EXCLUDED.employee_id, timesheet.employee_id)
            """, (emp_id, r["employee_name"], r["tab_number"], r["year"], r["month"], r["day"], r["mark"], r["hours"]))
            ok += 1
        except Exception as e:
            err += 1
            errors.append(str(e)[:100])
    return ok, err, "; ".join(errors[:5])


def import_orders(cur, file_bytes: bytes):
    orders, brigades, lines = parse_orders(file_bytes)
    ok, err = 0, 0
    errors = []

    # order_number -> id
    order_id_map = {}
    # brigade_id_local -> id
    brigade_id_map = {}

    for o in orders:
        try:
            cur.execute("""
                INSERT INTO orders (order_number, order_date, master_name, approver_name, approver_position, status)
                VALUES (%s, %s, %s, %s, %s, %s)
                ON CONFLICT (order_number) DO UPDATE SET
                    order_date = EXCLUDED.order_date,
                    master_name = EXCLUDED.master_name,
                    approver_name = EXCLUDED.approver_name,
                    status = EXCLUDED.status,
                    updated_at = NOW()
                RETURNING id
            """, (o["order_number"], o.get("order_date"), o["master_name"], o.get("approver_name"), o.get("approver_position"), o["status"]))
            oid = cur.fetchone()[0]
            order_id_map[o["order_number"]] = oid
            ok += 1
        except Exception as e:
            err += 1
            errors.append(str(e)[:100])

    for b in brigades:
        oid = order_id_map.get(b["order_number"])
        if not oid:
            continue
        try:
            cur.execute("""
                INSERT INTO order_brigades (order_id, brigade_number)
                VALUES (%s, %s)
                RETURNING id
            """, (oid, b["brigade_number"]))
            bid = cur.fetchone()[0]
            brigade_id_map[b["brigade_id_local"]] = bid
        except Exception as e:
            errors.append(str(e)[:100])

    for line in lines:
        oid = order_id_map.get(line["order_number"])
        bid = brigade_id_map.get(line["brigade_id_local"])
        if not oid:
            continue

        # Ищем employee_id по имени
        emp_id = None
        name = line["employee_name"]
        if name:
            cur.execute("SELECT id FROM employees WHERE full_name ILIKE %s LIMIT 1", (f"%{name}%",))
            row = cur.fetchone()
            if row:
                emp_id = row[0]

        try:
            cur.execute("""
                INSERT INTO order_lines
                    (order_id, brigade_id, employee_id, employee_name, position,
                     shift1_start, shift1_end, shift2_start, shift2_end,
                     hours_plan, hours_fact, is_ok)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            """, (
                oid, bid, emp_id, name, line.get("position"),
                line.get("shift1_start"), line.get("shift1_end"),
                line.get("shift2_start"), line.get("shift2_end"),
                line.get("hours_plan"), line.get("hours_fact"),
                line.get("is_ok", False),
            ))
            ok += 1
        except Exception as e:
            err += 1
            errors.append(str(e)[:100])

    return ok, err, "; ".join(errors[:5])


def debug_file(file_bytes: bytes) -> dict:
    """Возвращает структуру Excel-файла: листы и первые 15 строк каждого"""
    from openpyxl import load_workbook
    from io import BytesIO

    wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    result = {}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        sheet_rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 15:
                break
            sheet_rows.append([str(c) if c is not None else "" for c in row])
        result[sheet_name] = sheet_rows

    return {
        "statusCode": 200,
        "headers": CORS,
        "body": json.dumps({"sheets": list(wb.sheetnames), "preview": result}, ensure_ascii=False),
    }


def get_import_log(event):
    conn = get_conn()
    cur = conn.cursor()
    cur.execute("""
        SELECT id, file_name, file_type, rows_imported, rows_error, status, error_msg, imported_at
        FROM import_log ORDER BY imported_at DESC LIMIT 50
    """)
    rows = cur.fetchall()
    cur.close()
    conn.close()
    result = []
    for r in rows:
        result.append({
            "id": r[0], "file_name": r[1], "file_type": r[2],
            "rows_imported": r[3], "rows_error": r[4],
            "status": r[5], "error_msg": r[6],
            "imported_at": r[7].isoformat() if r[7] else None,
        })
    return {"statusCode": 200, "headers": CORS, "body": json.dumps(result)}


def get_data(event):
    """GET /excel-import?type=employees&limit=50"""
    params = event.get("queryStringParameters") or {}
    data_type = params.get("type", "employees")
    limit = int(params.get("limit", 100))

    conn = get_conn()
    cur = conn.cursor()

    if data_type == "employees":
        cur.execute("SELECT id, tab_number, full_name, position, category, work_schedule, status FROM employees ORDER BY category, full_name LIMIT %s", (limit,))
        cols = ["id", "tab_number", "full_name", "position", "category", "work_schedule", "status"]
    elif data_type == "work_types":
        cur.execute("SELECT id, code, name, group_name, unit, norm_per_hour FROM work_types ORDER BY group_name, name LIMIT %s", (limit,))
        cols = ["id", "code", "name", "group_name", "unit", "norm_per_hour"]
    elif data_type == "orders":
        cur.execute("SELECT id, order_number, order_date, master_name, status FROM orders ORDER BY order_date DESC LIMIT %s", (limit,))
        cols = ["id", "order_number", "order_date", "master_name", "status"]
    elif data_type == "timesheet":
        year = params.get("year", 2026)
        month = params.get("month", 5)
        cur.execute("SELECT employee_name, tab_number, day, mark, hours FROM timesheet WHERE year=%s AND month=%s ORDER BY employee_name, day LIMIT %s", (year, month, limit))
        cols = ["employee_name", "tab_number", "day", "mark", "hours"]
    else:
        cur.close()
        conn.close()
        return {"statusCode": 400, "headers": CORS, "body": json.dumps({"error": "unknown type"})}

    rows = cur.fetchall()
    cur.close()
    conn.close()

    result = []
    for r in rows:
        item = {}
        for i, col in enumerate(cols):
            v = r[i]
            if hasattr(v, "isoformat"):
                v = v.isoformat()
            item[col] = v
        result.append(item)

    return {"statusCode": 200, "headers": CORS, "body": json.dumps(result, ensure_ascii=False)}